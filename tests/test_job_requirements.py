from pathlib import Path

from openpyxl import Workbook

from resume_screening.job_requirements import apply_job_profile_overrides, load_job_requirements


def make_job_book(path: Path) -> None:
    workbook = Workbook()
    template = workbook.active
    template.title = "模板"
    template["A1"] = "岗位说明书"
    product = workbook.create_sheet("产品总监")
    product.append(["岗位说明书"])
    product.append([None, None, None, None, None, None, None, "HR(03)20260001"])
    product.append(["一、基本信息"])
    product.append(["岗位名称", "技术产品总监", "所属部门", "研发部", "直接上级", "郑廉", "岗位等级", None])
    product.append(["岗位编制", "1/1", "岗位性质", "全职", None, None, "职级", "总监"])
    product.append(["工作地点", "北京", "工作方式", "半线上", None, None, "工作经验", "8年以上"])
    product.append(["学历", "统招本科及以上", "专业", "计算机、信息工程", "证书持有", "", "技能/素质", "AI产品、Agent"])
    product.append(["语言要求", "普通话", None, None, "晋升路径", "", None, None])
    product.append(["二、岗位目的"])
    product.append(["具体描述", "负责AI产品规划和跨团队落地。"])
    product.append(["三、主要工作职责及核心考核指标"])
    product.append(["1. 核心职责", "负责产品路线图和核心指标。"])
    product.append(["2. 核心职责", "推动研发、算法、运营协同。"])
    product.append(["3. 核心职责", "建立产品数据反馈机制。"])
    product.append(["4. 补充职责", "有Agent经验优先。"])
    product.append(["四、其他说明"])
    product.append(["补充说明", "无"])
    empty = workbook.create_sheet("算法")
    empty.append(["岗位说明书"])
    empty.append(["一、基本信息"])
    workbook.save(path)


def test_load_job_requirements_ignores_template_and_detects_completeness(tmp_path: Path) -> None:
    path = tmp_path / "jobs.xlsx"
    make_job_book(path)

    jobs = load_job_requirements(path)

    assert "模板" not in jobs
    assert jobs["产品总监"].is_complete is True
    assert "负责AI产品规划" in jobs["产品总监"].raw_text
    assert jobs["算法"].is_complete is False
    assert "岗位名称" in jobs["算法"].missing_fields


def test_load_job_requirements_supports_tabular_job_rows(tmp_path: Path) -> None:
    path = tmp_path / "jobs.xlsx"
    workbook = Workbook()
    template = workbook.active
    template.title = "模板"
    sheet = workbook.create_sheet("后端开发")
    sheet.append(["岗位名称", "岗位别名", "岗位性质", "学历要求", "工作年限要求", "工作经验要求", "核心职责1", "核心职责2", "岗位jd"])
    sheet.append([
        "后端开发",
        "Java开发",
        "实习",
        "本科、硕士、博士",
        "26、27届毕业生",
        "0到1年",
        "Java后端开发",
        "Spring Cloud / Spring Boot",
        "负责医者核心业务后端开发。",
    ])
    sheet.append([
        "AI产品经理",
        "AI交互产品经理",
        "实习或全职",
        "本科、硕士、博士",
        "26、27届毕业生",
        "0-2年",
        "海外TOC产品",
        "AI交互设计",
        "负责ToC端AI交互产品规划。",
    ])
    workbook.save(path)

    jobs = load_job_requirements(path)

    assert jobs["后端开发"].is_complete is True
    assert jobs["后端开发"].sheet_name == "后端开发"
    assert jobs["后端开发"].fields["岗位别名"] == "Java开发"
    assert "Spring Cloud" in jobs["后端开发"].raw_text
    profile_lines = [line for line in jobs["后端开发"].fields["岗位核心画像"].splitlines() if line.strip()]
    assert 3 <= len(profile_lines) <= 10
    assert any(line.startswith("学历:") for line in profile_lines)
    assert any(line.startswith("工作内容:") for line in profile_lines)
    assert any(line.startswith("匹配程度:") for line in profile_lines)
    assert jobs["AI产品经理"].is_complete is True
    assert "AI交互产品经理" in jobs["AI产品经理"].raw_text


def test_load_job_requirements_supports_docx_core_profile(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "后端开发工程师岗位说明书.docx"
    document = Document()
    document.add_paragraph("岗位名称：后端开发工程师")
    document.add_paragraph("学历要求：本科及以上")
    document.add_paragraph("工作经验：3年以上 Java 后端开发经验")
    document.add_paragraph("岗位职责")
    document.add_paragraph("负责核心业务系统后端开发，参与架构设计和性能优化。")
    document.add_paragraph("任职要求")
    document.add_paragraph("熟悉 Spring Boot、MySQL、Redis，有高并发项目经验。")
    document.save(path)

    jobs = load_job_requirements(path)

    assert list(jobs) == ["后端开发工程师"]
    job = jobs["后端开发工程师"]
    assert job.is_complete is True
    assert job.fields["来源格式"] == "docx"
    assert "岗位核心画像" in job.fields
    profile_lines = [line for line in job.fields["岗位核心画像"].splitlines() if line.strip()]
    assert 3 <= len(profile_lines) <= 10
    assert any(line.startswith("学历:") for line in profile_lines)
    assert any(line.startswith("工作内容:") and "后端开发" in line for line in profile_lines)
    assert any(line.startswith("匹配程度:") for line in profile_lines)
    assert "Spring Boot" in job.fields["岗位核心画像"]
    assert "【岗位核心画像】" in job.raw_text


def test_docx_core_profile_summarizes_responsibilities_as_keywords_not_hard_checklist(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "产品运营岗位说明书.docx"
    document = Document()
    document.add_paragraph("岗位名称：AI产品运营")
    document.add_paragraph("岗位职责")
    document.add_paragraph("负责用户增长、活动运营、内容运营、数据分析和跨团队协同。")
    document.add_paragraph("负责社群维护、用户访谈、需求反馈、竞品调研和运营流程优化。")
    document.add_paragraph("任职要求")
    document.add_paragraph("熟悉AI产品，有运营数据分析经验，能推动产品和运营协作。")
    document.save(path)

    job = load_job_requirements(path)["AI产品运营"]
    profile = job.fields["岗位核心画像"]
    profile_lines = [line for line in profile.splitlines() if line.strip()]

    assert len(profile_lines) <= 10
    assert any(line.startswith("工作内容:") for line in profile_lines)
    assert any(line.startswith("匹配程度:") for line in profile_lines)


def test_load_job_requirements_supports_multiple_docx_job_blocks(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "岗位说明书.docx"
    document = Document()
    document.add_paragraph("岗位名称：后端开发工程师")
    document.add_paragraph("岗位职责：负责 Java 后端开发和服务稳定性建设。")
    document.add_paragraph("任职要求：本科及以上，熟悉 Spring Boot。")
    document.add_paragraph("岗位名称：AI产品经理")
    document.add_paragraph("岗位职责：负责 AI 产品规划、需求分析和跨团队推进。")
    document.add_paragraph("任职要求：熟悉 Agent 产品，有 ToC 产品经验。")
    document.save(path)

    jobs = load_job_requirements(path)

    assert set(jobs) == {"后端开发工程师", "AI产品经理"}
    assert jobs["AI产品经理"].is_complete is True
    assert 3 <= len([line for line in jobs["AI产品经理"].fields["岗位核心画像"].splitlines() if line.strip()]) <= 10
    assert "Agent" in jobs["AI产品经理"].fields["岗位核心画像"]


def test_apply_job_profile_overrides_adds_custom_only_job() -> None:
    jobs = {}

    updated = apply_job_profile_overrides(jobs, {"自定义增长岗位": "学历: 本科\n工作内容: 用户增长\n匹配程度: 核心相关"})

    assert "自定义增长岗位" in updated
    assert updated["自定义增长岗位"].fields["岗位名称"] == "自定义增长岗位"
    assert updated["自定义增长岗位"].fields["岗位核心画像"].startswith("学历: 本科")
    assert updated["自定义增长岗位"].is_complete is True


def test_apply_job_profile_overrides_compacts_legacy_long_profile() -> None:
    jobs = {}
    legacy_profile = "\n".join(
        [
            "岗位名称: 后端开发",
            "学历要求: 本科、硕士、博士",
            "年龄要求: 27岁以内",
            "性别要求: 不限",
            "核心职责1: Java后端开发",
            "核心职责2: Spring Cloud / Spring Boot",
            "核心职责3: 数据库设计优化",
            "岗位jd: 负责医者核心业务系统设计、开发和维护，确保高性能、高可用性和可扩展性。",
        ]
    )

    updated = apply_job_profile_overrides(jobs, {"后端开发": legacy_profile})
    profile_lines = [line for line in updated["后端开发"].fields["岗位核心画像"].splitlines() if line.strip()]

    assert 3 <= len(profile_lines) <= 10
    assert any(line.startswith("学历:") for line in profile_lines)
    assert any(line.startswith("年龄:") for line in profile_lines)
    assert any(line.startswith("性别:") for line in profile_lines)
    assert any(line.startswith("工作内容:") for line in profile_lines)
