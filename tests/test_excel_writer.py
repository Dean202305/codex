from pathlib import Path

from openpyxl import Workbook, load_workbook

from resume_screening.excel_writer import ResultWorkbookWriter
from resume_screening.models import DuplicateMatch, ParsedFilename, ScreeningResult, ScreeningScores


def create_result_book(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "多维表格"
    sheet.append(["ID", "当前阶段", "候选人", "附件", "岗位方向", "来源渠道", "候选人摘要（妙记）", "量化打分", "可投入周期", "面试负责人", "下次动作日", "不推进原因", "推荐/触达人"])
    sheet.append([1, "未通过", "历史候选人", "old.pdf", "财务", "BOSS", "", "", "", "", "", "不匹配", ""])
    workbook.save(path)


def make_result() -> ScreeningResult:
    return ScreeningResult(
        category="推荐初试",
        overall_score=8.5,
        summary="候选人具备财务管理经验。",
        screening_reason="经验匹配。",
        missing_information=[],
        reject_reason="",
        scores=ScreeningScores(8, 6, 8, 8, 9, 9, 8, 7),
    )


def test_writer_appends_columns_and_duplicate_red_fill(tmp_path: Path) -> None:
    path = tmp_path / "result.xlsx"
    create_result_book(path)
    writer = ResultWorkbookWriter(path)
    parsed = ParsedFilename(
        path=Path("【财务总监_北京 18-28K】郭燕婷 10年以上.pdf"),
        is_standard=True,
        job_name="财务总监",
        expected_location="北京",
        salary_range="18-28K",
        candidate_name="郭燕婷",
        work_experience="10年以上",
    )

    row_id = writer.append_result(
        parsed=parsed,
        result=make_result(),
        duplicate=DuplicateMatch(True, "old.pdf", 1, 1.0),
        extraction_errors=[],
        default_source_channel="",
        default_interviewer="",
    )
    writer.save()

    workbook = load_workbook(path)
    sheet = workbook["多维表格"]
    headers = [cell.value for cell in sheet[1]]
    assert headers[-3:] == ["初筛分类", "初筛理由", "缺失信息"]
    assert row_id == 2
    assert sheet["A3"].value == 2
    assert sheet["C3"].value == "郭燕婷"
    assert sheet["N3"].value == "推荐初试"
    assert "疑似重复简历" in sheet["P3"].value
    assert sheet["A3"].fill.fgColor.rgb == "FFFFC7CE"
