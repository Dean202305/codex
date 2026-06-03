from pathlib import Path

from openpyxl import Workbook

from resume_screening.config import AppConfig
from resume_screening.web.precheck import run_precheck


def create_result_book(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "多维表格"
    sheet.append(["ID", "当前阶段", "候选人", "附件", "岗位方向", "来源渠道", "候选人摘要（妙记）", "量化打分", "可投入周期", "面试负责人", "下次动作日", "不推进原因", "推荐/触达人"])
    workbook.save(path)


def create_job_book(path: Path) -> None:
    workbook = Workbook()
    template = workbook.active
    template.title = "模板"
    job = workbook.create_sheet("财务总监")
    job.append(["岗位名称", "财务总监", "学历", "本科"])
    job.append(["工作经验", "8年以上", "具体描述", "负责公司财务管理"])
    job.append(["1. 核心职责", "预算、核算、风控"])
    workbook.save(path)


def make_config(tmp_path: Path) -> AppConfig:
    resume_dir = tmp_path / "resumes"
    nested = resume_dir / "nested"
    nested.mkdir(parents=True)
    (resume_dir / "【财务总监_北京 18-28K】郭燕婷 10年以上.pdf").write_text("fake", encoding="utf-8")
    (nested / "【财务总监_上海 20-30K】李四 8年.txt").write_text("fake", encoding="utf-8")
    job_book = tmp_path / "jobs.xlsx"
    result_book = tmp_path / "result.xlsx"
    create_job_book(job_book)
    create_result_book(result_book)
    return AppConfig.model_validate(
        {
            "resume_dir": resume_dir,
            "job_book": job_book,
            "result_book": result_book,
            "index_path": tmp_path / "processed_index.json",
            "model": {"provider": "openai-compatible", "base_url": "", "api_key": "", "model": "", "allow_without_model": True},
        }
    )


def test_precheck_passes_for_valid_local_files(tmp_path: Path) -> None:
    result = run_precheck(make_config(tmp_path))

    assert result.status == "pass"
    assert result.resume_file_count == 2
    assert any(item.name == "岗位说明书" and item.status == "pass" for item in result.items)


def test_precheck_fails_when_result_book_is_missing(tmp_path: Path) -> None:
    config = make_config(tmp_path)
    config.result_book.unlink()

    result = run_precheck(config)

    assert result.status == "fail"
    assert any(item.name == "招聘结果表" and item.status == "fail" for item in result.items)


def test_precheck_warns_when_resume_jobs_do_not_match_job_sheets(tmp_path: Path) -> None:
    config = make_config(tmp_path)
    for file_path in config.resume_dir.rglob("*"):
        if file_path.is_file():
            file_path.unlink()
    (config.resume_dir / "后端开发工程师_北京_12-18K_张松_3年.pdf").write_text("fake", encoding="utf-8")

    result = run_precheck(config)

    assert result.status == "warning"
    assert any(
        item.name == "岗位匹配" and item.status == "warning" and "后端开发工程师" in item.message
        for item in result.items
    )


def test_precheck_warns_when_local_qwen_model_needs_download(tmp_path: Path, monkeypatch) -> None:
    config = make_config(tmp_path)
    runtime = tmp_path / "runtime" / "macos-arm64" / "llama-server"
    runtime.parent.mkdir(parents=True)
    runtime.write_text("#!/bin/sh\n", encoding="utf-8")
    monkeypatch.setenv("RESUME_SCREENING_RUNTIME_DIR", str(tmp_path / "runtime"))
    config.model = config.model.model_validate(
        {
            "provider": "local-qwen",
            "base_url": "",
            "api_key": "",
            "model": "qwen3.5-local",
            "timeout_seconds": 120,
            "temperature": 0.1,
            "allow_without_model": False,
            "local": {
                "model_path": str(tmp_path / "models" / "qwen.gguf"),
                "manifest_path": str(tmp_path / "models" / "manifest.json"),
                "host": "127.0.0.1",
                "port": 18080,
            },
        }
    )

    result = run_precheck(config)

    assert result.status == "warning"
    assert any(item.name == "本地模型" and item.status == "warning" and "下载" in item.message for item in result.items)


def test_precheck_passes_when_local_qwen_model_is_ready(tmp_path: Path, monkeypatch) -> None:
    config = make_config(tmp_path)
    runtime = tmp_path / "runtime" / "macos-arm64" / "llama-server"
    runtime.parent.mkdir(parents=True)
    runtime.write_text("#!/bin/sh\n", encoding="utf-8")
    model = tmp_path / "models" / "qwen.gguf"
    model.parent.mkdir()
    model.write_bytes(b"model")
    monkeypatch.setenv("RESUME_SCREENING_RUNTIME_DIR", str(tmp_path / "runtime"))
    config.model = config.model.model_validate(
        {
            "provider": "local-qwen",
            "base_url": "",
            "api_key": "",
            "model": "qwen3.5-local",
            "timeout_seconds": 120,
            "temperature": 0.1,
            "allow_without_model": False,
            "local": {
                "model_path": str(model),
                "manifest_path": str(tmp_path / "models" / "manifest.json"),
                "host": "127.0.0.1",
                "port": 18080,
            },
        }
    )

    result = run_precheck(config)

    assert result.status == "pass"
    assert any(item.name == "本地模型" and item.status == "pass" and "已就绪" in item.message for item in result.items)
