from pathlib import Path

from docx import Document
from openpyxl import Workbook

from resume_screening.config import AppConfig
from resume_screening.models import PipelineEvent
from resume_screening.pipeline import ScreeningPipeline


def create_result_book(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "多维表格"
    sheet.append(["ID", "当前阶段", "候选人", "附件", "岗位方向", "来源渠道", "候选人摘要（妙记）", "量化打分", "可投入周期", "面试负责人", "下次动作日", "不推进原因", "推荐/触达人"])
    workbook.save(path)


def create_incomplete_job_book(path: Path) -> None:
    workbook = Workbook()
    template = workbook.active
    template.title = "模板"
    job = workbook.create_sheet("财务总监")
    job.append(["岗位说明书"])
    job.append(["一、基本信息"])
    workbook.save(path)


def test_pipeline_emits_run_and_file_events(tmp_path: Path) -> None:
    resume_dir = tmp_path / "resumes"
    resume_dir.mkdir()
    resume = resume_dir / "【财务总监_北京 18-28K】郭燕婷 10年以上.docx"
    document = Document()
    document.add_paragraph("郭燕婷，10年以上财务经验。")
    document.save(resume)
    job_book = tmp_path / "jobs.xlsx"
    result_book = tmp_path / "result.xlsx"
    create_incomplete_job_book(job_book)
    create_result_book(result_book)
    events: list[PipelineEvent] = []
    config = AppConfig.model_validate(
        {
            "resume_dir": resume_dir,
            "job_book": job_book,
            "result_book": result_book,
            "index_path": tmp_path / "processed_index.json",
            "model": {"provider": "openai-compatible", "base_url": "", "api_key": "", "model": "", "allow_without_model": True},
        }
    )

    stats = ScreeningPipeline(config, event_handler=events.append).run()

    assert stats.processed == 1
    assert [event.type for event in events] == ["run_started", "file_started", "file_completed", "run_completed"]
    assert events[1].filename == resume.name
    assert events[-1].stats is stats
