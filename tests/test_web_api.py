from pathlib import Path
from threading import Event

from fastapi.testclient import TestClient
import httpx
from openpyxl import Workbook

from resume_screening.models import PipelineEvent, PipelineStats
from resume_screening.web.app import create_app
from resume_screening.web.runs import RunManager


def create_result_book(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "多维表格"
    sheet.append(["ID", "当前阶段", "候选人", "附件", "岗位方向", "来源渠道", "候选人摘要（妙记）", "量化打分", "可投入周期", "面试负责人", "下次动作日", "不推进原因", "推荐/触达人"])
    workbook.save(path)


def create_job_book(path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "模板"
    job = workbook.create_sheet("财务总监")
    job.append(["岗位名称", "财务总监", "学历", "本科"])
    job.append(["工作经验", "8年以上", "具体描述", "负责公司财务管理"])
    job.append(["1. 核心职责", "预算、核算、风控"])
    workbook.save(path)


def valid_payload(tmp_path: Path) -> dict:
    resume_dir = tmp_path / "resumes"
    nested = resume_dir / "nested"
    nested.mkdir(parents=True)
    (resume_dir / "【财务总监_北京 18-28K】郭燕婷 10年以上.pdf").write_text("fake", encoding="utf-8")
    (nested / "【财务总监_上海 20-30K】李四 8年.txt").write_text("fake", encoding="utf-8")
    job_book = tmp_path / "jobs.xlsx"
    result_book = tmp_path / "result.xlsx"
    create_job_book(job_book)
    create_result_book(result_book)
    return {
        "resume_dir": str(resume_dir),
        "job_book": str(job_book),
        "result_book": str(result_book),
        "index_path": str(tmp_path / "processed_index.json"),
        "ocr_command": "tesseract",
        "model": {"provider": "openai-compatible", "base_url": "", "api_key": "", "model": "", "timeout_seconds": 60, "temperature": 0.1, "allow_without_model": True},
    }


def test_config_endpoints_round_trip(tmp_path: Path) -> None:
    config_path = tmp_path / "config.yaml"
    client = TestClient(create_app(config_path=config_path))

    response = client.post("/api/config", json=valid_payload(tmp_path))
    assert response.status_code == 200
    assert response.json()["config"]["resume_dir"].endswith("resumes")

    response = client.get("/api/config")
    assert response.status_code == 200
    assert response.json()["config"]["result_book"].endswith("result.xlsx")


def test_precheck_endpoint_returns_items(tmp_path: Path) -> None:
    config_path = tmp_path / "config.yaml"
    client = TestClient(create_app(config_path=config_path))
    client.post("/api/config", json=valid_payload(tmp_path))

    response = client.get("/api/precheck")

    assert response.status_code == 200
    body = response.json()
    assert body["status"] in {"pass", "warning"}
    assert body["resume_file_count"] == 2
    assert body["items"]


def test_model_check_switches_unavailable_custom_model_to_local(tmp_path: Path, monkeypatch) -> None:
    def unavailable(*args, **kwargs):
        raise httpx.ConnectError("connection refused")

    monkeypatch.setattr("resume_screening.web.model_check.httpx.get", unavailable)
    config_path = tmp_path / "config.yaml"
    client = TestClient(create_app(config_path=config_path))
    payload = valid_payload(tmp_path)
    payload["model"] = {
        "provider": "openai-compatible",
        "base_url": "https://api.example.invalid/v1",
        "api_key": "test-key",
        "model": "custom-model",
        "timeout_seconds": 60,
        "temperature": 0.1,
        "allow_without_model": False,
        "fallback_to_local_when_unavailable": True,
    }
    client.post("/api/config", json=payload)

    response = client.post("/api/model/check")

    assert response.status_code == 200
    body = response.json()
    assert body["switched"] is True
    assert body["config"]["model"]["provider"] == "local-qwen"
    assert body["config"]["model"]["model"] == "qwen3.5-local"
    assert "自定义模型不可用" in body["message"]


def test_model_check_keeps_custom_model_when_local_fallback_disabled(tmp_path: Path, monkeypatch) -> None:
    def unavailable(*args, **kwargs):
        raise httpx.ConnectError("connection refused")

    monkeypatch.setattr("resume_screening.web.model_check.httpx.get", unavailable)
    config_path = tmp_path / "config.yaml"
    client = TestClient(create_app(config_path=config_path))
    payload = valid_payload(tmp_path)
    payload["model"] = {
        "provider": "openai-compatible",
        "base_url": "https://api.example.invalid/v1",
        "api_key": "test-key",
        "model": "custom-model",
        "timeout_seconds": 60,
        "temperature": 0.1,
        "allow_without_model": False,
        "fallback_to_local_when_unavailable": False,
    }
    client.post("/api/config", json=payload)

    response = client.post("/api/model/check")

    assert response.status_code == 200
    body = response.json()
    assert body["available"] is False
    assert body["switched"] is False
    assert body["config"]["model"]["provider"] == "openai-compatible"
    assert "自定义模型不可用" in body["message"]


def test_job_profiles_endpoint_returns_and_saves_overrides(tmp_path: Path) -> None:
    config_path = tmp_path / "config.yaml"
    client = TestClient(create_app(config_path=config_path))
    client.post("/api/config", json=valid_payload(tmp_path))

    response = client.get("/api/job-profiles")

    assert response.status_code == 200
    profiles = response.json()["profiles"]
    assert profiles[0]["job_name"] == "财务总监"
    assert "预算" in profiles[0]["profile"]

    save_response = client.post(
        "/api/job-profiles",
        json={"profiles": [{"job_name": "财务总监", "profile": "自定义画像：预算、风控、团队管理"}]},
    )

    assert save_response.status_code == 200
    saved_profile = save_response.json()["config"]["job_profile_overrides"]["财务总监"]
    assert "工作内容: 预算、风控、团队管理" in saved_profile
    assert len([line for line in saved_profile.splitlines() if line.strip()]) <= 10
    refreshed = client.get("/api/job-profiles").json()["profiles"]
    assert refreshed[0]["profile"] == saved_profile


def test_job_profiles_endpoint_saves_new_custom_job_profile(tmp_path: Path) -> None:
    config_path = tmp_path / "config.yaml"
    client = TestClient(create_app(config_path=config_path))
    client.post("/api/config", json=valid_payload(tmp_path))

    response = client.post(
        "/api/job-profiles",
        json={
            "profiles": [
                {"job_name": "财务总监", "profile": "学历: 本科\n工作内容: 预算、风控\n匹配程度: 核心相关"},
                {"job_name": "自定义运营岗位", "profile": "学历: 本科\n年龄: 30岁以内\n性别: 不限\n工作内容: 用户增长\n匹配程度: 核心相关"},
            ]
        },
    )

    assert response.status_code == 200
    names = [profile["job_name"] for profile in response.json()["profiles"]]
    assert "财务总监" in names
    assert "自定义运营岗位" in names
    assert response.json()["config"]["job_profile_overrides"]["自定义运营岗位"].startswith("学历: 本科")


def test_job_profiles_endpoint_prefers_model_generated_profile(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(
        "resume_screening.web.job_profiles.ModelClient.generate_job_profile",
        lambda self, job: f"模型画像：{job.sheet_name} 预算、风控、团队管理",
    )
    config_path = tmp_path / "config.yaml"
    client = TestClient(create_app(config_path=config_path))
    payload = valid_payload(tmp_path)
    payload["model"] = {
        "provider": "openai-compatible",
        "base_url": "https://api.example.com/v1",
        "api_key": "test-key",
        "model": "custom-model",
        "timeout_seconds": 60,
        "temperature": 0.1,
        "allow_without_model": False,
    }
    client.post("/api/config", json=payload)

    response = client.get("/api/job-profiles")

    assert response.status_code == 200
    profile = response.json()["profiles"][0]
    assert profile["profile"] == "模型画像：财务总监 预算、风控、团队管理"
    assert profile["source"] == "model"


def test_job_profiles_endpoint_uses_rule_profile_for_local_model(tmp_path: Path, monkeypatch) -> None:
    def fail_if_called(self, job):
        raise AssertionError("local model should not generate precheck job profiles")

    monkeypatch.setattr("resume_screening.web.job_profiles.ModelClient.generate_job_profile", fail_if_called)
    config_path = tmp_path / "config.yaml"
    client = TestClient(create_app(config_path=config_path))
    payload = valid_payload(tmp_path)
    payload["model"] = {
        "provider": "local-qwen",
        "base_url": "",
        "api_key": "",
        "model": "qwen3.5-local",
        "timeout_seconds": 120,
        "temperature": 0.1,
        "allow_without_model": False,
        "fallback_to_local_when_unavailable": True,
    }
    client.post("/api/config", json=payload)

    response = client.get("/api/job-profiles")

    assert response.status_code == 200
    profile = response.json()["profiles"][0]
    assert profile["source"] == "generated"
    assert profile["profile_error"] == ""
    assert "工作内容:" in profile["profile"]


def test_cancel_run_endpoint_marks_run_stopping(tmp_path: Path) -> None:
    started = Event()
    release = Event()

    class BlockingPipeline:
        def __init__(self, event_handler, should_stop) -> None:
            self.event_handler = event_handler
            self.should_stop = should_stop

        def run(self) -> PipelineStats:
            self.event_handler(PipelineEvent("run_started", "开始处理", total=1))
            started.set()
            release.wait(timeout=2)
            return PipelineStats()

    config_path = tmp_path / "config.yaml"
    client = TestClient(create_app(config_path=config_path, run_manager=RunManager(config_path, pipeline_factory=lambda config, handler, should_stop: BlockingPipeline(handler, should_stop))))
    client.post("/api/config", json=valid_payload(tmp_path))
    started_response = client.post("/api/runs")
    run_id = started_response.json()["run"]["run_id"]
    assert started.wait(timeout=2)

    response = client.post(f"/api/runs/{run_id}/cancel")
    release.set()

    assert response.status_code == 200
    assert response.json()["run"]["state"] == "stopping"
