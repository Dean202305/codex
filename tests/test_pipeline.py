from pathlib import Path

from openpyxl import Workbook, load_workbook

from resume_screening.config import AppConfig
from resume_screening.models import ScreeningResult, ScreeningScores
from resume_screening.pipeline import ScreeningPipeline


def create_result_book(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "多维表格"
    sheet.append(["ID", "当前阶段", "候选人", "附件", "岗位方向", "来源渠道", "候选人摘要（妙记）", "量化打分", "可投入周期", "面试负责人", "下次动作日", "不推进原因", "推荐/触达人"])
    workbook.save(path)


def create_incomplete_job_book(path: Path) -> None:
    workbook = Workbook()
    template = workbook.active
    template.title = "模板"
    template.append(["岗位说明书"])
    job = workbook.create_sheet("财务总监")
    job.append(["岗位说明书"])
    job.append(["一、基本信息"])
    workbook.save(path)


def create_complete_job_book(path: Path) -> None:
    workbook = Workbook()
    template = workbook.active
    template.title = "模板"
    job = workbook.create_sheet("财务总监")
    job.append(["岗位名称", "财务总监", "学历", "本科"])
    job.append(["工作经验", "8年以上", "具体描述", "负责财务管理"])
    job.append(["1. 核心职责", "预算、核算、风控"])
    workbook.save(path)


def test_pipeline_writes_manual_review_for_incomplete_job(tmp_path: Path) -> None:
    resume_dir = tmp_path / "resumes"
    resume_dir.mkdir()
    resume = resume_dir / "【财务总监_北京 18-28K】郭燕婷 10年以上.docx"
    from docx import Document
    document = Document()
    document.add_paragraph("郭燕婷，10年以上财务经验。")
    document.save(resume)
    job_book = tmp_path / "jobs.xlsx"
    result_book = tmp_path / "result.xlsx"
    create_incomplete_job_book(job_book)
    create_result_book(result_book)
    config = AppConfig.model_validate(
        {
            "resume_dir": resume_dir,
            "job_book": job_book,
            "result_book": result_book,
            "index_path": tmp_path / "processed_index.json",
            "model": {
                "provider": "openai-compatible",
                "base_url": "",
                "api_key": "",
                "model": "",
                "timeout_seconds": 60,
                "temperature": 0.1,
                "allow_without_model": True,
            },
            "screening": {
                "score_pass": 8,
                "score_excellent": 9,
                "categories": {
                    "recommend": "推荐初试",
                    "consider": "可考虑",
                    "manual": "待人工二筛",
                    "reject": "未通过",
                },
            },
        }
    )

    stats = ScreeningPipeline(config).run()

    assert stats.processed == 1
    assert stats.manual == 1
    workbook = load_workbook(result_book)
    sheet = workbook["多维表格"]
    assert sheet["B2"].value == "待人工二筛"
    assert "岗位要求不完整" in sheet["P2"].value


def test_pipeline_processes_nested_unsupported_file_as_manual_review(tmp_path: Path) -> None:
    resume_dir = tmp_path / "resumes"
    nested = resume_dir / "subfolder"
    nested.mkdir(parents=True)
    resume = nested / "【财务总监_北京 18-28K】郭燕婷 10年以上.txt"
    resume.write_text("郭燕婷，10年以上财务经验。", encoding="utf-8")
    job_book = tmp_path / "jobs.xlsx"
    result_book = tmp_path / "result.xlsx"
    create_incomplete_job_book(job_book)
    create_result_book(result_book)
    config = AppConfig.model_validate(
        {
            "resume_dir": resume_dir,
            "job_book": job_book,
            "result_book": result_book,
            "index_path": tmp_path / "processed_index.json",
            "model": {
                "provider": "openai-compatible",
                "base_url": "",
                "api_key": "",
                "model": "",
                "timeout_seconds": 60,
                "temperature": 0.1,
                "allow_without_model": True,
            },
        }
    )

    stats = ScreeningPipeline(config).run()

    assert stats.processed == 1
    assert stats.manual == 1
    workbook = load_workbook(result_book)
    sheet = workbook["多维表格"]
    assert sheet["D2"].value == resume.name
    assert sheet["B2"].value == "待人工二筛"
    assert "不支持的文件格式" in sheet["P2"].value


def test_pipeline_uses_model_when_config_is_complete_even_if_manual_fallback_allowed(tmp_path: Path, monkeypatch) -> None:
    resume_dir = tmp_path / "resumes"
    resume_dir.mkdir()
    resume = resume_dir / "【财务总监_北京 18-28K】郭燕婷 10年以上.docx"
    from docx import Document

    document = Document()
    document.add_paragraph("郭燕婷，10年以上财务经验。")
    document.save(resume)
    job_book = tmp_path / "jobs.xlsx"
    result_book = tmp_path / "result.xlsx"
    create_complete_job_book(job_book)
    create_result_book(result_book)

    class FakeModelClient:
        def __init__(self, config) -> None:
            self.config = config

        def evaluate(self, resume_text, job, filename_metadata) -> ScreeningResult:
            return ScreeningResult(
                category="推荐初试",
                overall_score=8.5,
                summary="模型已评估",
                screening_reason="岗位匹配",
                missing_information=[],
                reject_reason="",
                scores=ScreeningScores(8, 6, 8, 8, 9, 9, 8, 7),
            )

    monkeypatch.setattr("resume_screening.pipeline.ModelClient", FakeModelClient)
    config = AppConfig.model_validate(
        {
            "resume_dir": resume_dir,
            "job_book": job_book,
            "result_book": result_book,
            "index_path": tmp_path / "processed_index.json",
            "model": {
                "provider": "openai-compatible",
                "base_url": "https://api.example.com/v1",
                "api_key": "test-key",
                "model": "screening-model",
                "timeout_seconds": 60,
                "temperature": 0.1,
                "allow_without_model": True,
            },
        }
    )

    stats = ScreeningPipeline(config).run()

    assert stats.processed == 1
    assert stats.recommend == 1
    assert stats.manual == 0
    workbook = load_workbook(result_book)
    sheet = workbook["多维表格"]
    assert sheet["B2"].value == "推荐初试"
    assert sheet["G2"].value == "模型已评估"


def test_pipeline_uses_local_qwen_model_without_api_credentials(tmp_path: Path, monkeypatch) -> None:
    resume_dir = tmp_path / "resumes"
    resume_dir.mkdir()
    resume = resume_dir / "【财务总监_北京 18-28K】郭燕婷 10年以上.docx"
    from docx import Document

    document = Document()
    document.add_paragraph("郭燕婷，10年以上财务经验。")
    document.save(resume)
    job_book = tmp_path / "jobs.xlsx"
    result_book = tmp_path / "result.xlsx"
    create_complete_job_book(job_book)
    create_result_book(result_book)

    class FakeModelClient:
        def __init__(self, config) -> None:
            assert config.provider == "local-qwen"
            assert config.api_key == ""

        def evaluate(self, resume_text, job, filename_metadata) -> ScreeningResult:
            return ScreeningResult(
                category="推荐初试",
                overall_score=8.6,
                summary="本地模型已评估",
                screening_reason="候选人与岗位匹配。",
                missing_information=[],
                reject_reason="",
                scores=ScreeningScores(8, 6, 8, 8, 9, 9, 8, 7),
            )

    monkeypatch.setattr("resume_screening.pipeline.ModelClient", FakeModelClient)
    config = AppConfig.model_validate(
        {
            "resume_dir": resume_dir,
            "job_book": job_book,
            "result_book": result_book,
            "index_path": tmp_path / "processed_index.json",
            "model": {
                "provider": "local-qwen",
                "base_url": "",
                "api_key": "",
                "model": "qwen3.5-local",
                "timeout_seconds": 120,
                "temperature": 0.1,
                "allow_without_model": False,
                "local": {
                    "model_path": str(tmp_path / "models" / "qwen.gguf"),
                    "manifest_path": str(tmp_path / "models" / "manifest.json"),
                    "port": 18080,
                },
            },
        }
    )

    stats = ScreeningPipeline(config).run()

    assert stats.processed == 1
    assert stats.recommend == 1
    assert stats.manual == 0
    workbook = load_workbook(result_book)
    sheet = workbook["多维表格"]
    assert sheet["B2"].value == "推荐初试"
    assert sheet["G2"].value == "本地模型已评估"


def test_pipeline_uses_model_for_incomplete_job_when_model_is_configured(tmp_path: Path, monkeypatch) -> None:
    resume_dir = tmp_path / "resumes"
    resume_dir.mkdir()
    resume = resume_dir / "财务总监_北京_18-28K_郭燕婷_10年以上.docx"
    from docx import Document

    document = Document()
    document.add_paragraph("郭燕婷，10年以上财务经验。")
    document.save(resume)
    job_book = tmp_path / "jobs.xlsx"
    result_book = tmp_path / "result.xlsx"
    create_incomplete_job_book(job_book)
    create_result_book(result_book)

    class FakeModelClient:
        def __init__(self, config) -> None:
            self.config = config

        def evaluate(self, resume_text, job, filename_metadata) -> ScreeningResult:
            assert job.sheet_name == "财务总监"
            assert filename_metadata["job_name"] == "财务总监"
            return ScreeningResult(
                category="可考虑",
                overall_score=7.5,
                summary="模型已按现有岗位信息评估",
                screening_reason="岗位信息不完整，但候选人经验有一定匹配度。",
                missing_information=["模型认为仍需补充管理半径"],
                reject_reason="",
                scores=ScreeningScores(8, 6, 7, 7, 7.5, 8, 7, 6),
            )

    monkeypatch.setattr("resume_screening.pipeline.ModelClient", FakeModelClient)
    config = AppConfig.model_validate(
        {
            "resume_dir": resume_dir,
            "job_book": job_book,
            "result_book": result_book,
            "index_path": tmp_path / "processed_index.json",
            "model": {
                "provider": "openai-compatible",
                "base_url": "https://api.example.com/v1",
                "api_key": "test-key",
                "model": "screening-model",
                "timeout_seconds": 60,
                "temperature": 0.1,
                "allow_without_model": False,
            },
        }
    )

    stats = ScreeningPipeline(config).run()

    assert stats.processed == 1
    assert stats.consider == 1
    assert stats.manual == 0
    workbook = load_workbook(result_book)
    sheet = workbook["多维表格"]
    assert sheet["B2"].value == "可考虑"
    assert sheet["G2"].value == "模型已按现有岗位信息评估"
    assert "岗位要求不完整" in sheet["P2"].value
    assert "模型认为仍需补充管理半径" in sheet["P2"].value


def test_pipeline_uses_configured_job_alias_for_model_evaluation(tmp_path: Path, monkeypatch) -> None:
    resume_dir = tmp_path / "resumes"
    resume_dir.mkdir()
    resume = resume_dir / "后端开发工程师_北京_12-18K_张松_3年.docx"
    from docx import Document

    document = Document()
    document.add_paragraph("张松，3年后端开发经验。")
    document.save(resume)
    job_book = tmp_path / "jobs.xlsx"
    result_book = tmp_path / "result.xlsx"
    workbook = Workbook()
    template = workbook.active
    template.title = "模板"
    job = workbook.create_sheet("全栈")
    job.append(["岗位名称", "全栈工程师", "学历", "本科"])
    job.append(["工作经验", "3年以上", "具体描述", "负责后端和前端开发"])
    job.append(["1. 核心职责", "服务端开发、接口设计、系统稳定性"])
    workbook.save(job_book)
    create_result_book(result_book)

    class FakeModelClient:
        def __init__(self, config) -> None:
            self.config = config

        def evaluate(self, resume_text, job, filename_metadata) -> ScreeningResult:
            assert job.sheet_name == "全栈"
            assert filename_metadata["job_name"] == "后端开发工程师"
            assert filename_metadata["matched_job_sheet"] == "全栈"
            return ScreeningResult(
                category="推荐初试",
                overall_score=8.2,
                summary="模型已按别名岗位评估",
                screening_reason="后端经验与全栈岗位服务端要求匹配。",
                missing_information=[],
                reject_reason="",
                scores=ScreeningScores(8, 6, 8, 8, 8, 8, 8, 7),
            )

    monkeypatch.setattr("resume_screening.pipeline.ModelClient", FakeModelClient)
    config = AppConfig.model_validate(
        {
            "resume_dir": resume_dir,
            "job_book": job_book,
            "result_book": result_book,
            "index_path": tmp_path / "processed_index.json",
            "job_aliases": {"后端开发工程师": "全栈"},
            "model": {
                "provider": "openai-compatible",
                "base_url": "https://api.example.com/v1",
                "api_key": "test-key",
                "model": "screening-model",
                "timeout_seconds": 60,
                "temperature": 0.1,
                "allow_without_model": False,
            },
        }
    )

    stats = ScreeningPipeline(config).run()

    assert stats.recommend == 1
    workbook = load_workbook(result_book)
    sheet = workbook["多维表格"]
    assert sheet["E2"].value == "后端开发工程师"
    assert "岗位别名映射：后端开发工程师 -> 全栈" in sheet["P2"].value
