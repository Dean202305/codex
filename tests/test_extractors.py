from pathlib import Path
import subprocess
from types import SimpleNamespace

from resume_screening.extractors import extract_text


def test_extract_unsupported_file_returns_error(tmp_path: Path) -> None:
    path = tmp_path / "resume.txt"
    path.write_text("hello", encoding="utf-8")

    result = extract_text(path, ocr_command="tesseract")

    assert result.text == ""
    assert result.method == "unsupported"
    assert "不支持的文件格式" in result.errors


def test_extract_docx_reads_paragraphs(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("郭燕婷")
    document.add_paragraph("财务总监候选人，10年以上经验")
    document.save(path)

    result = extract_text(path, ocr_command="tesseract")

    assert "郭燕婷" in result.text
    assert "10年以上经验" in result.text
    assert result.method == "docx"


def test_short_pdf_text_is_not_treated_as_usable(monkeypatch, tmp_path: Path) -> None:
    class FakePage:
        def extract_text(self) -> str:
            return "1"

    class FakePdfReader:
        def __init__(self, path: str) -> None:
            self.pages = [FakePage()]

    monkeypatch.setattr("resume_screening.extractors.PdfReader", FakePdfReader)
    path = tmp_path / "resume.pdf"
    path.write_bytes(b"%PDF-1.4")

    result = extract_text(path, ocr_command="tesseract")

    assert result.text == ""
    assert result.method == "pdf"
    assert result.errors == ["PDF文本过短，可能需要OCR人工复核"]


def test_ocr_missing_command_returns_actionable_ocr_error(monkeypatch, tmp_path: Path) -> None:
    def raise_missing(*args, **kwargs):
        raise FileNotFoundError

    monkeypatch.setattr("resume_screening.extractors.subprocess.run", raise_missing)
    path = tmp_path / "resume.png"
    path.write_bytes(b"image")

    result = extract_text(path, ocr_command="missing-tesseract")

    assert result.text == ""
    assert result.method == "ocr"
    assert result.errors == ["OCR命令不存在：missing-tesseract"]


def test_ocr_timeout_returns_actionable_ocr_error(monkeypatch, tmp_path: Path) -> None:
    def raise_timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired(cmd="tesseract", timeout=120)

    monkeypatch.setattr("resume_screening.extractors.subprocess.run", raise_timeout)
    path = tmp_path / "resume.jpg"
    path.write_bytes(b"image")

    result = extract_text(path, ocr_command="tesseract")

    assert result.text == ""
    assert result.method == "ocr"
    assert result.errors == ["OCR超时"]


def test_ocr_nonzero_exit_returns_ocr_failure(monkeypatch, tmp_path: Path) -> None:
    def fake_run(*args, **kwargs):
        return SimpleNamespace(returncode=1, stdout="", stderr="bad image")

    monkeypatch.setattr("resume_screening.extractors.subprocess.run", fake_run)
    path = tmp_path / "resume.jpeg"
    path.write_bytes(b"image")

    result = extract_text(path, ocr_command="tesseract")

    assert result.text == ""
    assert result.method == "ocr"
    assert result.errors == ["OCR失败：bad image"]


def test_extract_docx_reads_table_cells(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "table_resume.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "候选人"
    table.cell(0, 1).text = "郭燕婷"
    document.save(path)

    result = extract_text(path, ocr_command="tesseract")

    assert "候选人" in result.text
    assert "郭燕婷" in result.text
    assert result.method == "docx"


def test_extract_pptx_reads_slide_text(tmp_path: Path) -> None:
    from pptx import Presentation

    path = tmp_path / "resume.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "郭燕婷"
    textbox = slide.shapes.add_textbox(0, 0, 3000000, 1000000)
    textbox.text = "财务总监候选人"
    presentation.save(path)

    result = extract_text(path, ocr_command="tesseract")

    assert "郭燕婷" in result.text
    assert "财务总监候选人" in result.text
    assert result.method == "pptx"
