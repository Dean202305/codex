from __future__ import annotations

import subprocess
from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from resume_screening.models import ExtractionResult


SUPPORTED_EXTENSIONS = {".pdf", ".jpg", ".jpeg", ".png", ".doc", ".docx", ".pptx"}


def extract_text(path: Path, ocr_command: str) -> ExtractionResult:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        return ExtractionResult(text="", method="unsupported", errors=["不支持的文件格式"])
    try:
        if suffix == ".pdf":
            return _extract_pdf(path)
        if suffix in {".jpg", ".jpeg", ".png"}:
            return _extract_image_ocr(path, ocr_command)
        if suffix == ".docx":
            return _extract_docx(path)
        if suffix == ".doc":
            return _extract_doc(path)
        if suffix == ".pptx":
            return _extract_pptx(path)
    except Exception as exc:
        return ExtractionResult(text="", method=suffix.lstrip("."), errors=[f"正文提取失败：{exc}"])
    return ExtractionResult(text="", method="unsupported", errors=["不支持的文件格式"])


def _extract_pdf(path: Path) -> ExtractionResult:
    reader = PdfReader(str(path))
    text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if len(text) < 80:
        return ExtractionResult(text=text, method="pdf", errors=["PDF文本过短，可能需要OCR人工复核"])
    return ExtractionResult(text=text, method="pdf")


def _extract_image_ocr(path: Path, ocr_command: str) -> ExtractionResult:
    completed = subprocess.run(
        [ocr_command, str(path), "stdout", "-l", "chi_sim+eng"],
        capture_output=True,
        text=True,
        check=False,
        timeout=120,
    )
    if completed.returncode != 0:
        return ExtractionResult(text="", method="ocr", errors=[f"OCR失败：{completed.stderr.strip()}"])
    return ExtractionResult(text=completed.stdout.strip(), method="ocr")


def _extract_docx(path: Path) -> ExtractionResult:
    document = Document(str(path))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    return ExtractionResult(text=text.strip(), method="docx")


def _extract_doc(path: Path) -> ExtractionResult:
    completed = subprocess.run(
        ["textutil", "-convert", "txt", "-stdout", str(path)],
        capture_output=True,
        text=True,
        check=False,
        timeout=120,
    )
    if completed.returncode != 0:
        return ExtractionResult(text="", method="doc", errors=[f"DOC转换失败：{completed.stderr.strip()}"])
    return ExtractionResult(text=completed.stdout.strip(), method="doc")


def _extract_pptx(path: Path) -> ExtractionResult:
    presentation = Presentation(str(path))
    parts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text.strip())
    return ExtractionResult(text="\n".join(parts), method="pptx")
