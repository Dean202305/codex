from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from resume_screening.models import CATEGORY_REJECT, DuplicateMatch, ParsedFilename, ScreeningResult


BASE_HEADERS = ["ID", "当前阶段", "候选人", "附件", "岗位方向", "来源渠道", "候选人摘要（妙记）", "量化打分", "可投入周期", "面试负责人", "下次动作日", "不推进原因", "推荐/触达人"]
EXTRA_HEADERS = ["初筛分类", "初筛理由", "缺失信息"]
DUPLICATE_FILL = PatternFill(fill_type="solid", fgColor="FFFFC7CE")


class ResultWorkbookWriter:
    def __init__(
        self,
        path: Path,
        sheet_name: str = "多维表格",
        *,
        score_pass: float = 7,
        score_excellent: float = 9,
    ) -> None:
        self.path = path
        self.score_pass = score_pass
        self.score_excellent = score_excellent
        self.workbook = load_workbook(path)
        self.sheet = self.workbook[sheet_name]
        self._ensure_headers()

    def _ensure_headers(self) -> None:
        headers = [cell.value for cell in self.sheet[1]]
        for header in EXTRA_HEADERS:
            if header not in headers:
                self.sheet.cell(row=1, column=len(headers) + 1).value = header
                headers.append(header)

    def next_id(self) -> int:
        ids: list[int] = []
        for row in self.sheet.iter_rows(min_row=2, min_col=1, max_col=1, values_only=True):
            value = row[0]
            if isinstance(value, int):
                ids.append(value)
        return (max(ids) if ids else 0) + 1

    def append_result(
        self,
        parsed: ParsedFilename,
        result: ScreeningResult,
        duplicate: DuplicateMatch,
        extraction_errors: list[str],
        default_source_channel: str,
        default_interviewer: str,
    ) -> int:
        row_id = self.next_id()
        missing = list(result.missing_information) + list(parsed.missing_fields) + list(extraction_errors)
        if duplicate.is_duplicate:
            missing.append(f"疑似重复简历：与历史附件 {duplicate.matched_filename} 内容一致")
        row = [
            row_id,
            result.category,
            parsed.candidate_name,
            parsed.path.name,
            parsed.job_name,
            default_source_channel,
            result.summary,
            format_score_block(result, score_pass=self.score_pass, score_excellent=self.score_excellent),
            "",
            default_interviewer,
            "",
            result.reject_reason if result.category == CATEGORY_REJECT else "",
            "",
            result.category,
            result.screening_reason,
            "；".join(item for item in missing if item),
        ]
        self.sheet.append(row)
        row_number = self.sheet.max_row
        if duplicate.is_duplicate:
            for cell in self.sheet[row_number]:
                cell.fill = DUPLICATE_FILL
        return row_id

    def save(self) -> None:
        self.workbook.save(self.path)


def format_score_block(result: ScreeningResult, *, score_pass: float = 7, score_excellent: float = 9) -> str:
    conclusion = "优秀" if result.overall_score >= score_excellent else "通过" if result.overall_score >= score_pass else "未通过"
    scores = result.scores
    return "\n".join(
        [
            f"综合评分：{result.overall_score:g}/10",
            f"结论：{conclusion}",
            "",
            f"能力：{scores.ability:g}",
            f"ego大小：{scores.ego:g}",
            f"野心desire：{scores.desire:g}",
            f"学习能力与聪明程度：{scores.learning_ability:g}",
            "",
            f"岗位匹配度：{scores.job_fit:g}",
            f"经验匹配：{scores.experience_fit:g}",
            f"技能匹配：{scores.skill_fit:g}",
            f"稳定性/风险：{scores.stability_risk:g}",
            "",
            f"评价：{result.screening_reason}",
        ]
    )
