from __future__ import annotations

from dataclasses import replace
from pathlib import Path
import re

from openpyxl import load_workbook

from resume_screening.extractors import extract_text
from resume_screening.models import JobRequirement


PLACEHOLDER_MARKERS = ("XXXX", "具体描述", "无则填", "年       月")
REQUIRED_FIELDS = ("岗位名称", "学历", "工作经验", "具体描述", "1. 核心职责")
TABULAR_REQUIRED_FIELDS = ("岗位名称",)
TABULAR_DETAIL_FIELDS = ("岗位jd", "核心职责1", "核心职责2", "核心职责3", "工作经验要求", "学历要求")
SPREADSHEET_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
DOCUMENT_EXTENSIONS = {".docx", ".doc"}
DOCUMENT_JOB_NAME_RE = re.compile(r"^(?:岗位名称|职位名称|招聘岗位|应聘岗位|岗位)[:：\s]+(.{2,60})$")
DOCUMENT_FIELD_RE = re.compile(r"^([^:：]{2,18})[:：]\s*(.+)$")
EDUCATION_RE = re.compile(r"(统招本科及以上|本科及以上|硕士及以上|博士|硕士|研究生|本科|大专及以上|大专|专科)")
AGE_RE = re.compile(r"(?:年龄要求|年龄)[:：\s]*([^\n，。；;]{1,24})|(\d{2}\s*(?:[-~至到]\s*\d{2})?\s*岁(?:以内|以下|以上|左右)?)")
GENDER_RE = re.compile(r"(?:性别要求|性别)[:：\s]*([^\n，。；;]{1,12})|(男性|女性|男士|女士|男|女|不限)")
CORE_PROFILE_KEYWORDS = (
    "岗位画像",
    "候选人画像",
    "核心画像",
    "岗位职责",
    "核心职责",
    "主要职责",
    "工作职责",
    "工作内容",
    "任职要求",
    "岗位要求",
    "任职资格",
    "能力要求",
    "技能要求",
    "工作经验",
    "项目经验",
    "学历",
    "专业",
    "加分",
    "优先",
)
PROFILE_BASIC_KEYS = (
    "岗位名称",
    "岗位别名",
    "工作地点",
    "工作方式",
    "岗位性质",
    "学历",
    "学历要求",
    "专业",
    "工作经验",
    "工作年限要求",
    "薪资",
)
PROFILE_EDUCATION_KEYS = ("学历", "学历要求", "最低学历", "教育背景")
PROFILE_AGE_KEYS = ("年龄", "年龄要求")
PROFILE_GENDER_KEYS = ("性别", "性别要求")
PROFILE_EXPERIENCE_KEYS = ("工作经验", "工作经验要求", "工作年限要求", "年限要求")
PROFILE_LOCATION_KEYS = ("工作地点", "地点")
DUTY_KEYWORDS = ("岗位职责", "核心职责", "主要职责", "工作职责", "工作内容", "职责")
REQUIREMENT_KEYWORDS = (
    "任职要求",
    "岗位要求",
    "任职资格",
    "能力要求",
    "技能要求",
    "工作经验",
    "项目经验",
    "学历",
    "专业",
)
BONUS_KEYWORDS = ("加分", "优先", "额外", "其他要求")
COMMON_PROFILE_PREFIXES = (
    "负责",
    "参与",
    "推动",
    "建立",
    "搭建",
    "制定",
    "熟悉",
    "具备",
    "具有",
    "掌握",
    "了解",
    "要求",
    "能够",
    "能",
    "有",
)


def _clean(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _is_placeholder(value: str) -> bool:
    return not value or any(marker in value for marker in PLACEHOLDER_MARKERS)


def _extract_sheet_fields(rows: list[list[str]]) -> dict[str, str]:
    fields: dict[str, str] = {}
    for row in rows:
        for index in range(0, len(row), 2):
            key = row[index].strip() if index < len(row) else ""
            value = row[index + 1].strip() if index + 1 < len(row) else ""
            if key and value:
                fields[key] = value
    return fields


def _raw_text(rows: list[list[str]]) -> str:
    lines = []
    for row in rows:
        text = " ".join(cell for cell in row if cell)
        if text:
            lines.append(text)
    return "\n".join(lines)


def _missing_required(fields: dict[str, str]) -> list[str]:
    missing = []
    for field_name in REQUIRED_FIELDS:
        value = fields.get(field_name, "")
        if _is_placeholder(value):
            missing.append(field_name)
    return missing


def _is_tabular_sheet(rows: list[list[str]]) -> bool:
    if not rows:
        return False
    header = rows[0]
    return "岗位名称" in header and any(name in header for name in ("岗位别名", "岗位jd", "核心职责1", "学历要求"))


def _tabular_missing_required(fields: dict[str, str]) -> list[str]:
    missing = [field_name for field_name in TABULAR_REQUIRED_FIELDS if _is_placeholder(fields.get(field_name, ""))]
    if not any(not _is_placeholder(fields.get(field_name, "")) for field_name in TABULAR_DETAIL_FIELDS):
        missing.append("岗位要求")
    return missing


def _row_raw_text(fields: dict[str, str]) -> str:
    return "\n".join(f"{key}: {value}" for key, value in fields.items() if value)


def _load_tabular_jobs(sheet_name: str, rows: list[list[str]]) -> dict[str, JobRequirement]:
    header = rows[0]
    jobs: dict[str, JobRequirement] = {}
    for row in rows[1:]:
        fields = {
            header[index]: row[index]
            for index in range(min(len(header), len(row)))
            if header[index] and row[index]
        }
        job_name = fields.get("岗位名称", "").strip()
        if not job_name:
            continue
        missing = _tabular_missing_required(fields)
        raw_text = _row_raw_text(fields)
        fields["岗位核心画像"] = _build_core_profile(raw_text, fields)
        unique_name = _unique_job_key(jobs, job_name, sheet_name)
        jobs[unique_name] = JobRequirement(
            sheet_name=sheet_name,
            fields=fields,
            raw_text=_document_raw_text(fields, raw_text),
            is_complete=len(missing) == 0,
            missing_fields=missing,
        )
    return jobs


def load_job_requirements(path: Path, *, ocr_command: str = "tesseract") -> dict[str, JobRequirement]:
    suffix = path.suffix.lower()
    if suffix in SPREADSHEET_EXTENSIONS:
        return _load_spreadsheet_job_requirements(path)
    if suffix in DOCUMENT_EXTENSIONS:
        return _load_document_job_requirements(path, ocr_command=ocr_command)
    supported = ", ".join(sorted(SPREADSHEET_EXTENSIONS | DOCUMENT_EXTENSIONS))
    raise ValueError(f"岗位说明书格式暂不支持：{suffix or '无扩展名'}，支持 {supported}")


def apply_job_profile_overrides(
    jobs: dict[str, JobRequirement],
    overrides: dict[str, str],
) -> dict[str, JobRequirement]:
    if not overrides:
        return jobs
    updated = dict(jobs)
    for job_name, profile in overrides.items():
        job_name = job_name.strip()
        cleaned_profile = profile.strip()
        if not cleaned_profile or not job_name:
            continue
        compacted_profile = compact_job_profile_text(cleaned_profile, job_name=job_name)
        if job_name not in updated:
            fields = {
                "岗位名称": job_name,
                "岗位核心画像": compacted_profile,
                "来源格式": "自定义岗位画像",
            }
            updated[job_name] = JobRequirement(
                sheet_name=job_name,
                fields=fields,
                raw_text=_document_raw_text(fields, ""),
                is_complete=True,
                missing_fields=[],
            )
            continue
        job = updated[job_name]
        fields = dict(job.fields)
        fields["岗位核心画像"] = compacted_profile
        missing = [field for field in job.missing_fields if field != "岗位核心画像"]
        updated[job_name] = replace(
            job,
            fields=fields,
            raw_text=_document_raw_text(fields, job.raw_text),
            is_complete=len(missing) == 0,
            missing_fields=missing,
        )
    return updated


def compact_job_profile_text(profile: str, *, job_name: str = "") -> str:
    compact_lines = [line.strip() for line in profile.splitlines() if line.strip()]
    if _is_compact_profile(compact_lines):
        return "\n".join(compact_lines[:10])
    fields = _extract_document_fields(profile)
    if job_name:
        fields.setdefault("岗位名称", job_name)
    compacted = _build_core_profile(profile, fields)
    return compacted or profile.strip()


def _is_compact_profile(lines: list[str]) -> bool:
    if not 3 <= len(lines) <= 10:
        return False
    labels = {line.split(":", 1)[0].strip() for line in lines if ":" in line}
    required = {"学历", "年龄", "性别", "工作内容", "匹配程度"}
    allowed = {*required, "经验", "技能", "加分", "地点", "关键词"}
    return required.issubset(labels) and labels.issubset(allowed)


def _load_spreadsheet_job_requirements(path: Path) -> dict[str, JobRequirement]:
    workbook = load_workbook(path, data_only=True)
    jobs: dict[str, JobRequirement] = {}
    for sheet_name in workbook.sheetnames:
        if sheet_name == "模板":
            continue
        sheet = workbook[sheet_name]
        rows = [[_clean(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
        if _is_tabular_sheet(rows):
            jobs.update(_load_tabular_jobs(sheet_name, rows))
            continue
        fields = _extract_sheet_fields(rows)
        missing = _missing_required(fields)
        raw_text = _raw_text(rows)
        fields["岗位核心画像"] = _build_core_profile(raw_text, fields)
        unique_name = _unique_job_key(jobs, sheet_name, sheet_name)
        jobs[unique_name] = JobRequirement(
            sheet_name=sheet_name,
            fields=fields,
            raw_text=_document_raw_text(fields, raw_text),
            is_complete=len(missing) == 0,
            missing_fields=missing,
        )
    return jobs


def _load_document_job_requirements(path: Path, *, ocr_command: str) -> dict[str, JobRequirement]:
    extraction = extract_text(path, ocr_command)
    if extraction.errors:
        raise ValueError("；".join(extraction.errors))
    raw_text = extraction.text.strip()
    if not raw_text:
        raise ValueError("岗位说明书正文为空")

    jobs: dict[str, JobRequirement] = {}
    for job_name, block in _split_document_job_blocks(raw_text, path):
        fields = _extract_document_fields(block)
        fields.setdefault("岗位名称", job_name)
        fields["来源格式"] = path.suffix.lower().lstrip(".")
        fields["岗位核心画像"] = _build_core_profile(block, fields)
        missing = _document_missing_required(fields)
        unique_name = _unique_job_key(jobs, job_name, path.stem)
        jobs[unique_name] = JobRequirement(
            sheet_name=unique_name,
            fields=fields,
            raw_text=_document_raw_text(fields, block),
            is_complete=len(missing) == 0,
            missing_fields=missing,
        )
    return jobs


def _split_document_job_blocks(raw_text: str, path: Path) -> list[tuple[str, str]]:
    lines = _document_lines(raw_text)
    matches: list[tuple[int, str]] = []
    for index, line in enumerate(lines):
        job_name = _document_job_name_from_line(line)
        if job_name:
            matches.append((index, job_name))
    if matches:
        blocks: list[tuple[str, str]] = []
        for position, (start, job_name) in enumerate(matches):
            end = matches[position + 1][0] if position + 1 < len(matches) else len(lines)
            blocks.append((job_name, "\n".join(lines[start:end])))
        return blocks

    fallback = _document_job_name_from_title(path.stem, lines)
    return [(fallback, raw_text)]


def _document_job_name_from_line(line: str) -> str:
    match = DOCUMENT_JOB_NAME_RE.match(line.strip())
    if not match:
        return ""
    return _clean_document_job_name(match.group(1))


def _document_job_name_from_title(stem: str, lines: list[str]) -> str:
    candidates = [stem, *(lines[:3])]
    for candidate in candidates:
        cleaned = _clean_document_job_name(candidate)
        if cleaned:
            return cleaned
    return stem.strip() or "岗位说明书"


def _unique_job_key(existing: dict[str, JobRequirement], job_name: str, source: str = "") -> str:
    if job_name not in existing:
        return job_name
    suffix = _clean_document_job_name(source) or "重复岗位"
    candidate = f"{job_name}（{suffix}）"
    counter = 2
    while candidate in existing:
        candidate = f"{job_name}（{suffix}{counter}）"
        counter += 1
    return candidate


def _clean_document_job_name(value: str) -> str:
    cleaned = re.sub(r"[\s_/-]+", " ", value).strip(" ：:-_")
    cleaned = re.sub(r"(岗位说明书|职位说明书|岗位JD|招聘需求|岗位需求|JD)$", "", cleaned, flags=re.IGNORECASE).strip()
    if not cleaned or cleaned in {"岗位说明书", "职位说明书", "招聘需求"}:
        return ""
    return cleaned[:60]


def _extract_document_fields(raw_text: str) -> dict[str, str]:
    fields: dict[str, str] = {}
    for line in _document_lines(raw_text):
        match = DOCUMENT_FIELD_RE.match(line)
        if not match:
            continue
        key = match.group(1).strip()
        value = match.group(2).strip()
        if key and value:
            fields[key] = value
    return fields


def _build_core_profile(raw_text: str, fields: dict[str, str]) -> str:
    lines = _document_lines(raw_text)
    selected: list[str] = []
    duty_items = _collect_profile_items(lines, DUTY_KEYWORDS, REQUIREMENT_KEYWORDS + BONUS_KEYWORDS)
    requirement_items = _collect_profile_items(lines, REQUIREMENT_KEYWORDS, DUTY_KEYWORDS + BONUS_KEYWORDS)
    bonus_items = _collect_profile_items(lines, BONUS_KEYWORDS, DUTY_KEYWORDS + REQUIREMENT_KEYWORDS)

    field_duties = [value for key, value in fields.items() if "职责" in key and value]
    if field_duties:
        duty_items.extend(field_duties)
    if not duty_items:
        duty_items = _keyword_nearby_lines(lines, DUTY_KEYWORDS)
    if not requirement_items:
        requirement_items = _keyword_nearby_lines(lines, REQUIREMENT_KEYWORDS)
    if not bonus_items:
        bonus_items = _keyword_nearby_lines(lines, BONUS_KEYWORDS)

    _add_profile_line(selected, "学历", _first_field_value(fields, PROFILE_EDUCATION_KEYS) or _regex_value(EDUCATION_RE, raw_text) or "未写明")
    _add_profile_line(selected, "年龄", _first_field_value(fields, PROFILE_AGE_KEYS) or _regex_value(AGE_RE, raw_text) or "未写明")
    _add_profile_line(selected, "性别", _first_field_value(fields, PROFILE_GENDER_KEYS) or _regex_value(GENDER_RE, raw_text) or "未写明")

    work_terms = _extract_terms([*duty_items, *field_duties], max_terms=3)
    if not work_terms:
        work_terms = _extract_terms(lines[:20], max_terms=3)
    work_terms = _usable_profile_terms(work_terms)
    _add_profile_line(selected, "工作内容", "、".join(work_terms) or _first_field_value(fields, ("岗位名称",)) or "未写明", max_chars=64)
    _add_profile_line(selected, "匹配程度", "核心工作内容匹配优先", max_chars=32)

    _add_profile_line(selected, "经验", _first_field_value(fields, PROFILE_EXPERIENCE_KEYS), max_chars=36)
    skill_terms = _profile_skill_terms(requirement_items, max_terms=3)
    _add_profile_line(selected, "技能", "、".join(skill_terms), max_chars=64)
    bonus_terms = _extract_terms(bonus_items, max_terms=2)
    _add_profile_line(selected, "加分", "、".join(bonus_terms), max_chars=48)
    _add_profile_line(selected, "地点", _first_field_value(fields, PROFILE_LOCATION_KEYS), max_chars=32)

    unique = list(dict.fromkeys(item for item in selected if item))
    return "\n".join(unique[:10])


def _add_profile_line(lines: list[str], label: str, value: str, *, max_chars: int = 40) -> None:
    cleaned = _short_profile_value(value, max_chars=max_chars)
    if cleaned:
        lines.append(f"{label}: {cleaned}")


def _first_field_value(fields: dict[str, str], keys: tuple[str, ...]) -> str:
    for key in keys:
        value = fields.get(key, "").strip()
        if value and not _is_placeholder(value):
            return value
    return ""


def _regex_value(pattern: re.Pattern[str], raw_text: str) -> str:
    match = pattern.search(raw_text)
    if not match:
        return ""
    for group in match.groups():
        if group:
            return group.strip()
    return match.group(0).strip()


def _short_profile_value(value: str, *, max_chars: int) -> str:
    cleaned = re.sub(r"\s+", " ", value).strip(" ：:，,。；;、")
    if not cleaned or _is_placeholder(cleaned):
        return ""
    if len(cleaned) <= max_chars:
        return cleaned
    return cleaned[:max_chars].rstrip(" ：:，,。；;、") + "..."


def _collect_profile_items(
    lines: list[str],
    start_keywords: tuple[str, ...],
    stop_keywords: tuple[str, ...],
) -> list[str]:
    items: list[str] = []
    active = False
    remaining = 0
    for line in lines:
        if active and _contains_any(line, stop_keywords):
            active = False
            remaining = 0
        if _contains_any(line, start_keywords):
            active = True
            remaining = 8
            value = _line_value_after_label(line)
            if value:
                items.append(value)
            elif not _looks_like_section_heading(line):
                items.append(line)
            continue
        if active and remaining > 0:
            if _looks_like_section_heading(line):
                active = False
                remaining = 0
                continue
            items.append(line)
            remaining -= 1
    return items


def _keyword_nearby_lines(lines: list[str], keywords: tuple[str, ...]) -> list[str]:
    indexes: set[int] = set()
    for index, line in enumerate(lines):
        if _contains_any(line, keywords):
            indexes.update(range(index, min(index + 5, len(lines))))
    return [lines[index] for index in sorted(indexes)]


def _format_profile_terms(label: str, items: list[str], *, max_terms: int) -> str:
    terms = _extract_terms(items, max_terms=max_terms)
    if not terms:
        return ""
    return f"{label}：{'、'.join(terms)}"


def _profile_skill_terms(items: list[str], *, max_terms: int) -> list[str]:
    terms = _extract_terms(items, max_terms=max_terms * 4)
    filtered = [
        term
        for term in terms
        if term not in {"以上", "及以上", "不限"} and not any(word in term for word in ("学历", "本科", "硕士", "博士", "大专", "专科"))
    ]
    return _usable_profile_terms(filtered or terms)[:max_terms]


def _usable_profile_terms(terms: list[str]) -> list[str]:
    usable: list[str] = []
    for term in terms:
        cleaned = _clean_profile_item(term)
        if (
            not cleaned
            or _is_placeholder(cleaned)
            or cleaned in {"岗位说明书", "补充职责", "补充说明"}
            or _looks_like_section_heading(cleaned)
        ):
            continue
        if cleaned not in usable:
            usable.append(cleaned)
    return usable


def _extract_terms(items: list[str], *, max_terms: int) -> list[str]:
    terms: list[str] = []
    for item in items:
        cleaned = _clean_profile_item(item)
        if not cleaned:
            continue
        for part in re.split(r"[、,，；;；/|]+", cleaned):
            term = _clean_profile_item(part)
            if not term:
                continue
            for splitter in ("以及", "并且", "和", "及", "与"):
                if splitter in term:
                    split_terms = [_clean_profile_item(piece) for piece in term.split(splitter)]
                    if all(split_terms):
                        for split_term in split_terms:
                            if split_term not in terms:
                                terms.append(split_term)
                        break
            else:
                if term not in terms:
                    terms.append(term)
            if len(terms) >= max_terms:
                return terms
    return terms


def _clean_profile_item(value: str) -> str:
    cleaned = re.sub(r"^[\s\-•·●○*（(]*\d*[）).、．]?\s*", "", value).strip()
    cleaned = _line_value_after_label(cleaned) or cleaned
    cleaned = cleaned.strip(" ：:，,。；;、")
    for prefix in COMMON_PROFILE_PREFIXES:
        if cleaned.startswith(prefix) and len(cleaned) > len(prefix) + 2:
            cleaned = cleaned[len(prefix) :].strip(" ：:，,。；;、")
            break
    if cleaned in CORE_PROFILE_KEYWORDS or _looks_like_section_heading(cleaned):
        return ""
    return cleaned[:60]


def _line_value_after_label(line: str) -> str:
    match = DOCUMENT_FIELD_RE.match(line.strip())
    if not match:
        return ""
    return match.group(2).strip()


def _contains_any(line: str, keywords: tuple[str, ...]) -> bool:
    return any(keyword in line for keyword in keywords)


def _looks_like_section_heading(line: str) -> bool:
    stripped = line.strip()
    if len(stripped) > 28:
        return False
    if re.match(r"^[一二三四五六七八九十\d]+[、.．]", stripped):
        return True
    return stripped in set(CORE_PROFILE_KEYWORDS)


def _document_missing_required(fields: dict[str, str]) -> list[str]:
    missing: list[str] = []
    if _is_placeholder(fields.get("岗位名称", "")):
        missing.append("岗位名称")
    if _is_placeholder(fields.get("岗位核心画像", "")) or len(fields.get("岗位核心画像", "")) < 40:
        missing.append("岗位核心画像")
    return missing


def _document_raw_text(fields: dict[str, str], raw_text: str) -> str:
    return "\n".join(
        [
            "【岗位核心画像】",
            fields.get("岗位核心画像", ""),
            "",
            "【岗位说明书原文】",
            raw_text,
        ]
    ).strip()


def _document_lines(raw_text: str) -> list[str]:
    return [line.strip() for line in raw_text.splitlines() if line.strip()]
